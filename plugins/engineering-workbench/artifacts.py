"""Draft artifact exporters; optional dependencies never install themselves."""
from __future__ import annotations

import csv
import hashlib
import importlib
import json
import math
import os
import shutil
import tempfile
from pathlib import Path
from xml.sax.saxutils import escape

from .core import VERSION, digest, inside, text

FORMATS = {"md", "csv", "docx", "xlsx", "pptx", "pdf"}
DRAFT = "DRAFT - sources and layout require review; not an approved engineering deliverable."


def dependency(name):
    try:
        return importlib.import_module(name)
    except ImportError:
        raise ValueError(f"Missing optional dependency: {name}; install requirements-office.txt in the Hermes Python environment") from None


def validate(args):
    kind = args.get("format")
    if kind not in FORMATS:
        raise ValueError("Unsupported format")
    title = text(args.get("title"), "title", 160)
    sections = args.get("sections", [])
    if not isinstance(sections, list) or len(sections) > 100:
        raise ValueError("sections must be a list with at most 100 entries")
    normalized = []
    for section in sections:
        if not isinstance(section, dict) or set(section) != {"heading", "body"}:
            raise ValueError("Each section needs heading and body")
        normalized.append({"heading": text(section["heading"], "heading", 160),
                           "body": text(section["body"], "body", 10000)})
    columns, rows = args.get("columns", []), args.get("rows", [])
    if not isinstance(columns, list) or len(columns) > 30:
        raise ValueError("columns must contain at most 30 names")
    columns = [text(c, "column", 100) for c in columns]
    if len(set(columns)) != len(columns):
        raise ValueError("Duplicate column names")
    if not isinstance(rows, list) or len(rows) > 1000:
        raise ValueError("rows must contain at most 1000 entries")
    for row in rows:
        if not isinstance(row, list) or len(row) != len(columns):
            raise ValueError("Every row must match columns")
        for cell in row:
            if not (cell is None or isinstance(cell, (str, int, float, bool))):
                raise ValueError("Cells must be scalar JSON values")
            if isinstance(cell, float) and not math.isfinite(cell):
                raise ValueError("Non-finite cell")
            if isinstance(cell, int) and abs(cell) > 10**30:
                raise ValueError("Integer cell too large")
            if isinstance(cell, str) and cell:
                text(cell, "cell", 5000)
    sources = args.get("sources")
    if not isinstance(sources, list) or not 1 <= len(sources) <= 100:
        raise ValueError("sources must contain 1..100 references, such as user-supplied data identifiers")
    sources = [text(s, "source", 2000) for s in sources]
    if kind in {"csv", "xlsx"} and not columns:
        raise ValueError("Tabular exports require columns")
    if kind not in {"csv", "xlsx"} and not normalized:
        raise ValueError("Document exports require sections")
    data = {"format": kind, "title": title, "sections": normalized, "columns": columns, "rows": rows, "sources": sources}
    if len(json.dumps(data, ensure_ascii=False)) > 250000:
        raise ValueError("Artifact payload exceeds 250000 characters")
    return data


def csv_cell(value):
    # Preserve numbers; neutralize user text that spreadsheet applications may execute.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def write_file(path, data):
    kind, title, sections = data["format"], data["title"], data["sections"]
    columns, rows, sources = data["columns"], data["rows"], data["sources"]
    if kind == "csv":
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow([csv_cell(c) for c in columns])
            writer.writerows([[csv_cell(c) for c in row] for row in rows])
    elif kind == "md":
        # Plain user content, never interpreted as commands by this exporter.
        parts = [f"# {title}", DRAFT]
        parts += [f"## {s['heading']}\n\n{s['body']}" for s in sections]
        if columns:
            parts.append("## Data (JSON)\n\n" + json.dumps({"columns": columns, "rows": rows}, ensure_ascii=False, indent=2))
        parts.append("## Source references (supplied, not independently verified)\n\n" + "\n".join(sources))
        path.write_text("\n\n".join(parts) + "\n", encoding="utf-8")
    elif kind == "docx":
        doc = dependency("docx").Document()
        doc.add_heading(title, 0)
        doc.add_paragraph(DRAFT)
        for s in sections:
            doc.add_heading(s["heading"], 1)
            doc.add_paragraph(s["body"])
        if columns:
            table = doc.add_table(rows=1, cols=len(columns))
            table.style = "Table Grid"
            for cell, label in zip(table.rows[0].cells, columns):
                cell.text = label
            for row in rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = "" if value is None else str(value)
        doc.add_heading("Source references (not independently verified)", 1)
        for source in sources:
            doc.add_paragraph(source)
        doc.save(path)
    elif kind == "xlsx":
        xl = dependency("openpyxl")
        styles = importlib.import_module("openpyxl.styles")
        book = xl.Workbook()
        sheet = book.active
        sheet.title = "Data"
        for ri, row in enumerate([columns] + rows, 1):
            for ci, value in enumerate(row, 1):
                if isinstance(value, int) and not isinstance(value, bool) and abs(value) >= 10**15:
                    value = str(value)  # Excel otherwise silently loses integer precision.
                cell = sheet.cell(ri, ci, value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Literal text, including =...; never an executable formula.
                cell.alignment = styles.Alignment(vertical="top", wrap_text=True)
                if ri == 1:
                    cell.font = styles.Font(bold=True)
        for ci in range(1, len(columns) + 1):
            sheet.column_dimensions[xl.utils.get_column_letter(ci)].width = 24
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        meta = book.create_sheet("Provenance")
        for row in [["Title", title], ["Status", DRAFT], ["Sources verified", False]] + [["Source", s] for s in sources]:
            meta.append(row)
            for cell in meta[meta.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        meta.column_dimensions["A"].width = 24
        meta.column_dimensions["B"].width = 90
        if sections:
            narrative = book.create_sheet("Narrative")
            for section in sections:
                narrative.append([section["heading"], section["body"]])
                for cell in narrative[narrative.max_row]:
                    cell.data_type = "s"
                    cell.alignment = styles.Alignment(vertical="top", wrap_text=True)
            narrative.column_dimensions["A"].width = 24
            narrative.column_dimensions["B"].width = 90
        book.save(path)
    elif kind == "pptx":
        pptx = dependency("pptx")
        from pptx.util import Inches, Pt
        deck = pptx.Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        def slide(heading, body):
            s = deck.slides.add_slide(deck.slide_layouts[1])
            s.shapes.title.text = heading
            s.placeholders[1].text = body
            for p in s.placeholders[1].text_frame.paragraphs:
                p.font.size = Pt(18)
        slide(title, DRAFT)
        import textwrap
        for s in sections:
            for part in textwrap.wrap(s["body"], width=650, replace_whitespace=False) or [""]:
                slide(s["heading"], part)
        if columns:
            # Avoid silently discarding tabular data in a presentation.
            raise ValueError("PPTX tables are not supported; export tables separately as XLSX or DOCX")
        for source in sources:
            for part in textwrap.wrap(source, width=650, replace_whitespace=False):
                slide("Source reference (not independently verified)", part)
        deck.save(path)
    else:
        dependency("reportlab")
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        if columns:
            raise ValueError("PDF tables are not supported; export a DOCX/XLSX and review its rendered PDF")
        font = "Helvetica"
        font_path = os.environ.get("HERMES_WORKBENCH_PDF_FONT")
        if font_path:
            font = "WorkbenchUnicode"
            pdfmetrics.registerFont(TTFont(font, font_path))
        elif any(ord(c) > 127 for c in title + str(sections) + str(sources)):
            raise ValueError("Unicode PDF needs HERMES_WORKBENCH_PDF_FONT pointing to a licensed Unicode TTF; no font is bundled")
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = font
        story = [Paragraph(escape(title), styles["Title"]), Paragraph(DRAFT, styles["Normal"]), Spacer(1, 12)]
        for s in sections:
            story += [Paragraph(escape(s["heading"]), styles["Heading2"]), Paragraph(escape(s["body"]).replace("\n", "<br/>"), styles["Normal"])]
        story += [Paragraph("Source references (not independently verified)", styles["Heading2"])]
        story += [Paragraph(escape(source), styles["Normal"]) for source in sources]
        SimpleDocTemplate(str(path)).build(story)


def export(args):
    data = validate(args)
    folder = inside("artifacts")
    folder.mkdir(mode=0o700, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix="workbench-", dir=folder))
    path = directory / ("draft." + data["format"])
    try:
        write_file(path, data)
        supplemental = []
        if data["format"] == "csv" and data["sections"]:
            narrative = directory / "narrative.json"
            narrative.write_text(json.dumps({"title": data["title"], "sections": data["sections"], "status": DRAFT}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            supplemental.append({"file": narrative.name, "path": str(narrative), "sha256": hashlib.sha256(narrative.read_bytes()).hexdigest()})
        evidence = {"version": VERSION, "supplemental_files": supplemental, "status": "draft_generated", "file": path.name,
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "input_sha256": digest(data),
                    "sources": data["sources"], "sources_verified": False, "visual_review": "not_performed",
                    "warning": DRAFT}
        manifest = directory / "manifest.json"
        manifest.write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return {**evidence, "path": str(path), "manifest": str(manifest)}
    except Exception:
        shutil.rmtree(directory)
        raise
